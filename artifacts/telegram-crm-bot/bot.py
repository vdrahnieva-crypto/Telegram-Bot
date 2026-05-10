import os
import io
import re
import logging
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import (
    Application,
    ApplicationHandlerStop,
    CommandHandler,
    CallbackQueryHandler,
    MessageHandler,
    filters,
    ContextTypes,
    ConversationHandler,
)
from database import Database, DB_PATH

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO
)
logger = logging.getLogger(__name__)

db = Database()

ADD_NAME, ADD_PHONE, ADD_NOTES, ADD_TAGS = range(4)
SEARCH_BY_NAME, SEARCH_BY_PHONE = range(2)
EDIT_VALUE = 0
BL_ADD_PHONE, BL_ADD_REASON = range(2)
BL_REMOVE_PHONE = 0
IMPORT_TEXT = 0


def main_menu_keyboard():
    keyboard = [
        [InlineKeyboardButton("➕ Добавить клиента", callback_data="add_client")],
        [InlineKeyboardButton("🔍 Найти клиента", callback_data="find_client")],
        [InlineKeyboardButton("📋 Все клиенты", callback_data="all_clients")],
        [InlineKeyboardButton("🚫 ЧС список", callback_data="blacklist_menu")],
        [InlineKeyboardButton("📥 Скачать Excel", callback_data="export_excel")],
        [InlineKeyboardButton("💾 Резервная копия", callback_data="backup_db")],
        [InlineKeyboardButton("📋 Импорт из текста", callback_data="import_text")],
    ]
    return InlineKeyboardMarkup(keyboard)


def find_menu_keyboard():
    keyboard = [
        [InlineKeyboardButton("👤 По имени", callback_data="search_by_name")],
        [InlineKeyboardButton("📞 По телефону", callback_data="search_by_phone")],
        [InlineKeyboardButton("🏷 По тегу", callback_data="search_by_tag")],
        [InlineKeyboardButton("❌ Отмена", callback_data="menu")],
    ]
    return InlineKeyboardMarkup(keyboard)


def blacklist_menu_keyboard():
    keyboard = [
        [InlineKeyboardButton("➕ Добавить в ЧС", callback_data="bl_add")],
        [InlineKeyboardButton("📋 Показать ЧС", callback_data="bl_show")],
        [InlineKeyboardButton("❌ Удалить из ЧС", callback_data="bl_remove")],
        [InlineKeyboardButton("⬅️ Назад", callback_data="menu")],
    ]
    return InlineKeyboardMarkup(keyboard)


def _build_add_tags_keyboard(selected_ids: list, all_tags: list) -> InlineKeyboardMarkup:
    rows = []
    row = []
    for tag_id, tag_name in all_tags:
        label = f"✓ {tag_name}" if tag_id in selected_ids else tag_name
        row.append(InlineKeyboardButton(label, callback_data=f"addtag_{tag_id}"))
        if len(row) == 2:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    rows.append([InlineKeyboardButton("✅ Готово", callback_data="addtags_done")])
    return InlineKeyboardMarkup(rows)


def _build_client_tags_keyboard(client_id: int, selected_ids: list, all_tags: list) -> InlineKeyboardMarkup:
    rows = []
    row = []
    for tag_id, tag_name in all_tags:
        label = f"✓ {tag_name}" if tag_id in selected_ids else tag_name
        row.append(InlineKeyboardButton(label, callback_data=f"ctag_{client_id}_{tag_id}"))
        if len(row) == 2:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    rows.append([InlineKeyboardButton("✅ Сохранить", callback_data=f"ctags_done_{client_id}")])
    rows.append([InlineKeyboardButton("❌ Отмена", callback_data="menu")])
    return InlineKeyboardMarkup(rows)


def _build_search_tag_keyboard(all_tags: list) -> InlineKeyboardMarkup:
    rows = []
    row = []
    for tag_id, tag_name in all_tags:
        row.append(InlineKeyboardButton(tag_name, callback_data=f"stag_{tag_id}"))
        if len(row) == 2:
            rows.append(row)
            row = []
    if row:
        rows.append(row)
    rows.append([InlineKeyboardButton("⬅️ Назад", callback_data="find_client")])
    return InlineKeyboardMarkup(rows)


_GROUP_ONLY_TEXT = "⛔ Бот работает только внутри группы."


async def _private_block_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_chat and update.effective_chat.type == "private":
        await update.message.reply_text(_GROUP_ONLY_TEXT)
        raise ApplicationHandlerStop


async def _private_block_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_chat and update.effective_chat.type == "private":
        await update.callback_query.answer(_GROUP_ONLY_TEXT, show_alert=True)
        raise ApplicationHandlerStop


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "👋 Добро пожаловать в CRM-бот!\n\nЧто хотите сделать?",
        reply_markup=main_menu_keyboard(),
    )


async def menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    await query.edit_message_text(
        "Что хотите сделать?",
        reply_markup=main_menu_keyboard(),
    )


async def add_client_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    context.user_data.clear()
    await query.edit_message_text(
        "📝 Добавляем нового клиента.\n\nВведите *полное имя* клиента:",
        parse_mode="Markdown"
    )
    return ADD_NAME


async def add_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["name"] = update.message.text.strip()
    await update.message.reply_text("📞 Введите *номер телефона* клиента:", parse_mode="Markdown")
    return ADD_PHONE


async def add_phone(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["phone"] = update.message.text.strip()
    await update.message.reply_text(
        "📝 Добавьте *заметку* о клиенте.\n\nОтправьте *Пропустить*, чтобы оставить пустой:",
        parse_mode="Markdown"
    )
    return ADD_NOTES


async def add_notes(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    context.user_data["notes"] = "" if text.lower() == "пропустить" else text
    return await _show_add_tags(update.message, context)


async def _show_add_tags(message, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["selected_tags"] = []
    all_tags = db.get_all_tags()
    keyboard = _build_add_tags_keyboard([], all_tags)
    await message.reply_text(
        "🏷 Выберите *теги* для клиента (можно несколько).\n\nНажмите *✅ Готово*, когда выберете нужные:",
        parse_mode="Markdown",
        reply_markup=keyboard,
    )
    return ADD_TAGS


async def toggle_add_tag(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    tag_id = int(query.data.split("_")[1])
    selected = context.user_data.setdefault("selected_tags", [])
    if tag_id in selected:
        selected.remove(tag_id)
    else:
        selected.append(tag_id)
    all_tags = db.get_all_tags()
    keyboard = _build_add_tags_keyboard(selected, all_tags)
    await query.edit_message_reply_markup(reply_markup=keyboard)
    return ADD_TAGS


async def finish_add_tags(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    data = context.user_data
    user = query.from_user
    client_id = db.add_client(
        name=data.get("name", ""),
        phone=data.get("phone", ""),
        email=data.get("email", ""),
        notes=data.get("notes", ""),
        added_by_user_id=user.id,
        added_by_username=user.username,
    )
    selected_tags = data.get("selected_tags", [])
    if selected_tags:
        db.set_client_tags(client_id, selected_tags)
    all_tags = db.get_all_tags()
    tags_map = {t[0]: t[1] for t in all_tags}
    tags_line = ", ".join(tags_map[t] for t in selected_tags if t in tags_map) or "—"
    keyboard = [[InlineKeyboardButton("⬅️ В меню", callback_data="menu")]]
    await query.edit_message_text(
        f"✅ *Клиент сохранён!*\n\n"
        f"👤 Имя: {data.get('name')}\n"
        f"📞 Телефон: {data.get('phone')}\n"
        f"📝 Заметка: {data.get('notes') or '—'}\n"
        f"🏷 Теги: {tags_line}\n"
        f"🆔 ID: `{client_id}`",
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )
    context.user_data.clear()
    return ConversationHandler.END


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text(
        "❌ Отменено.",
        reply_markup=main_menu_keyboard(),
    )
    return ConversationHandler.END


async def find_client_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    await query.edit_message_text(
        "🔍 Как хотите найти клиента?",
        reply_markup=find_menu_keyboard(),
    )


async def search_by_name_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    await query.edit_message_text("👤 Введите *имя* клиента для поиска:", parse_mode="Markdown")
    return SEARCH_BY_NAME


async def search_by_phone_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    await query.edit_message_text("📞 Введите *номер телефона* для поиска:", parse_mode="Markdown")
    return SEARCH_BY_PHONE


async def do_search_by_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query_text = update.message.text.strip()
    clients = db.search_by_name(query_text)
    await _send_search_results(update, clients, query_text)
    return ConversationHandler.END


async def do_search_by_phone(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query_text = update.message.text.strip()
    clients = db.search_by_phone(query_text)
    await _send_search_results(update, clients, query_text)
    return ConversationHandler.END


async def _send_search_results(update: Update, clients, query_text: str):
    if not clients:
        keyboard = [
            [InlineKeyboardButton("🔍 Найти снова", callback_data="find_client")],
            [InlineKeyboardButton("⬅️ В меню", callback_data="menu")],
        ]
        await update.message.reply_text(
            f"😕 Клиент не найден: *{query_text}*",
            parse_mode="Markdown",
            reply_markup=InlineKeyboardMarkup(keyboard),
        )
        return
    for client in clients:
        await send_client_card(update, client)
    await update.message.reply_text(
        f"Найдено клиентов: *{len(clients)}*.",
        parse_mode="Markdown",
        reply_markup=main_menu_keyboard(),
    )


async def search_by_tag_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    all_tags = db.get_all_tags()
    keyboard = _build_search_tag_keyboard(all_tags)
    await query.edit_message_text(
        "🏷 Выберите тег для поиска:",
        reply_markup=keyboard,
    )


async def do_search_by_tag(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    tag_id = int(query.data.split("_")[1])
    all_tags = db.get_all_tags()
    tag_name = next((t[1] for t in all_tags if t[0] == tag_id), "—")
    clients = db.get_clients_by_tag(tag_id)
    if not clients:
        keyboard = [
            [InlineKeyboardButton("🏷 Другой тег", callback_data="search_by_tag")],
            [InlineKeyboardButton("⬅️ В меню", callback_data="menu")],
        ]
        await query.edit_message_text(
            f"😕 Клиентов с тегом *{tag_name}* не найдено.",
            parse_mode="Markdown",
            reply_markup=InlineKeyboardMarkup(keyboard),
        )
        return
    await query.edit_message_text(
        f"🏷 Тег *{tag_name}* — найдено клиентов: {len(clients)}",
        parse_mode="Markdown",
    )
    for client in clients:
        await send_client_card(query, client)
    await query.message.reply_text("Что хотите сделать?", reply_markup=main_menu_keyboard())


async def send_client_card(update_or_query, client):
    cid, name, phone, email, notes, created_at, added_by_uid, added_by_uname = (*client, None, None)[:8]
    tags = db.get_client_tags(cid)
    tags_line = ", ".join(t[1] for t in tags) if tags else "—"
    keyboard = [
        [
            InlineKeyboardButton("✏️ Изменить", callback_data=f"edit_{cid}"),
            InlineKeyboardButton("🗑 Удалить", callback_data=f"delete_{cid}"),
        ],
        [InlineKeyboardButton("📝 Изменить заметку", callback_data=f"editfield_{cid}_notes")],
        [InlineKeyboardButton("🏷 Изменить теги", callback_data=f"client_tags_{cid}")],
    ]
    if added_by_uname:
        added_line = f"\n👤 Добавил: @{added_by_uname}"
    elif added_by_uid:
        added_line = f"\n👤 Добавил: #{added_by_uid}"
    else:
        added_line = ""
    text = (
        f"👤 *{name}*\n"
        f"📞 {phone}\n"
        f"📝 {notes or '—'}\n"
        f"🏷 {tags_line}\n"
        f"📅 Добавлен: {created_at[:10]}"
        f"{added_line}\n"
        f"🆔 ID: `{cid}`"
    )
    if hasattr(update_or_query, "message"):
        await update_or_query.message.reply_text(
            text, parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(keyboard)
        )
    else:
        await update_or_query.reply_text(
            text, parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(keyboard)
        )


async def client_tags_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    client_id = int(query.data.split("_")[2])
    client = db.get_client(client_id)
    if not client:
        await query.edit_message_text("Клиент не найден.")
        return
    all_tags = db.get_all_tags()
    current_tags = db.get_client_tags(client_id)
    selected_ids = [t[0] for t in current_tags]
    context.user_data["edit_selected_tags"] = selected_ids.copy()
    keyboard = _build_client_tags_keyboard(client_id, selected_ids, all_tags)
    await query.edit_message_text(
        f"🏷 Теги клиента *{client[1]}*:\n\nВыберите нужные теги:",
        parse_mode="Markdown",
        reply_markup=keyboard,
    )


async def toggle_client_tag(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    parts = query.data.split("_")
    client_id = int(parts[1])
    tag_id = int(parts[2])
    selected = context.user_data.get("edit_selected_tags", [])
    if tag_id in selected:
        selected.remove(tag_id)
    else:
        selected.append(tag_id)
    context.user_data["edit_selected_tags"] = selected
    all_tags = db.get_all_tags()
    keyboard = _build_client_tags_keyboard(client_id, selected, all_tags)
    await query.edit_message_reply_markup(reply_markup=keyboard)


async def finish_client_tags(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    client_id = int(query.data.split("_")[2])
    selected = context.user_data.get("edit_selected_tags", [])
    db.set_client_tags(client_id, selected)
    client = db.get_client(client_id)
    tags = db.get_client_tags(client_id)
    tags_line = ", ".join(t[1] for t in tags) if tags else "—"
    keyboard = [[InlineKeyboardButton("⬅️ В меню", callback_data="menu")]]
    await query.edit_message_text(
        f"✅ Теги обновлены!\n\n"
        f"👤 *{client[1]}*\n"
        f"🏷 {tags_line}",
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )
    context.user_data.pop("edit_selected_tags", None)


async def all_clients(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    clients = db.get_all_clients()
    if not clients:
        await query.edit_message_text(
            "📋 Клиентов пока нет. Добавьте первого!",
            reply_markup=main_menu_keyboard(),
        )
        return
    await query.edit_message_text(f"📋 Всего клиентов: *{len(clients)}*", parse_mode="Markdown")
    for client in clients:
        await send_client_card(query, client)
    await query.message.reply_text("Что хотите сделать?", reply_markup=main_menu_keyboard())


async def delete_client(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    client_id = int(query.data.split("_")[1])
    client = db.get_client(client_id)
    if not client:
        await query.edit_message_text("Клиент не найден.")
        return
    keyboard = [
        [
            InlineKeyboardButton("✅ Да, удалить", callback_data=f"confirm_delete_{client_id}"),
            InlineKeyboardButton("❌ Отмена", callback_data="menu"),
        ]
    ]
    await query.edit_message_text(
        f"🗑 Удалить клиента *{client[1]}*?",
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )


async def confirm_delete(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    client_id = int(query.data.split("_")[2])
    client = db.get_client(client_id)
    name = client[1] if client else "Клиент"
    db.delete_client(client_id)
    await query.edit_message_text(
        f"✅ *{name}* удалён.",
        parse_mode="Markdown",
        reply_markup=main_menu_keyboard(),
    )


async def edit_client(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    client_id = int(query.data.split("_")[1])
    context.user_data["edit_client_id"] = client_id
    client = db.get_client(client_id)
    if not client:
        await query.edit_message_text("Клиент не найден.")
        return
    keyboard = [
        [InlineKeyboardButton("👤 Имя", callback_data=f"editfield_{client_id}_name")],
        [InlineKeyboardButton("📞 Телефон", callback_data=f"editfield_{client_id}_phone")],
        [InlineKeyboardButton("📝 Заметка", callback_data=f"editfield_{client_id}_notes")],
        [InlineKeyboardButton("❌ Отмена", callback_data="menu")],
    ]
    await query.edit_message_text(
        f"✏️ Редактирование: *{client[1]}*\n\nКакое поле изменить?",
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )


async def edit_field_prompt(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    parts = query.data.split("_")
    client_id = int(parts[1])
    field = parts[2]
    context.user_data["edit_client_id"] = client_id
    context.user_data["edit_field"] = field
    field_labels = {"name": "Имя", "phone": "Телефон", "notes": "Заметка"}
    await query.edit_message_text(
        f"✏️ Введите новое значение для *{field_labels.get(field, field)}*:",
        parse_mode="Markdown"
    )
    return EDIT_VALUE


async def edit_value_save(update: Update, context: ContextTypes.DEFAULT_TYPE):
    new_value = update.message.text.strip()
    client_id = context.user_data.get("edit_client_id")
    field = context.user_data.get("edit_field")
    db.update_client_field(client_id, field, new_value)
    client = db.get_client(client_id)
    keyboard = [[InlineKeyboardButton("⬅️ В меню", callback_data="menu")]]
    field_labels = {"name": "Имя", "phone": "Телефон", "notes": "Заметка"}
    await update.message.reply_text(
        f"✅ *{field_labels.get(field, field)}* обновлено!\n\n"
        f"👤 Имя: {client[1]}\n"
        f"📞 Телефон: {client[2]}\n"
        f"📝 Заметка: {client[4] or '—'}",
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )
    context.user_data.clear()
    return ConversationHandler.END


async def blacklist_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    blacklist = db.get_blacklist()
    count = len(blacklist)
    await query.edit_message_text(
        f"🚫 *ЧС список* — {count} номер(ов)\n\nВыберите действие:",
        parse_mode="Markdown",
        reply_markup=blacklist_menu_keyboard(),
    )


async def bl_add_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    await query.edit_message_text(
        "🚫 Введите *номер телефона* для добавления в чёрный список:",
        parse_mode="Markdown",
    )
    return BL_ADD_PHONE


async def bl_add_phone(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["bl_phone"] = update.message.text.strip()
    await update.message.reply_text(
        "📝 Укажите *причину* (или /skip, чтобы пропустить):",
        parse_mode="Markdown",
    )
    return BL_ADD_REASON


async def bl_add_reason(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["bl_reason"] = update.message.text.strip()
    return await bl_save(update, context)


async def bl_skip_reason(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["bl_reason"] = ""
    return await bl_save(update, context)


async def bl_save(update: Update, context: ContextTypes.DEFAULT_TYPE):
    phone = context.user_data.get("bl_phone", "")
    reason = context.user_data.get("bl_reason", "")
    added = db.add_to_blacklist(phone, reason)
    keyboard = [[InlineKeyboardButton("⬅️ ЧС список", callback_data="blacklist_menu")]]
    if added:
        text = (
            f"✅ Номер *{phone}* добавлен в чёрный список.\n"
            f"📝 Причина: {reason or '—'}"
        )
    else:
        text = f"⚠️ Номер *{phone}* уже есть в чёрном списке."
    await update.message.reply_text(
        text,
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )
    context.user_data.clear()
    return ConversationHandler.END


async def bl_show(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    blacklist = db.get_blacklist()
    if not blacklist:
        await query.edit_message_text(
            "📋 Чёрный список пуст.",
            reply_markup=blacklist_menu_keyboard(),
        )
        return
    lines = []
    for i, (bid, phone, reason, created_at) in enumerate(blacklist, 1):
        entry = f"{i}. 📞 `{phone}`"
        if reason:
            entry += f"\n    📝 {reason}"
        entry += f"\n    📅 {created_at[:10]}"
        lines.append(entry)
    text = "🚫 *Чёрный список:*\n\n" + "\n\n".join(lines)
    await query.edit_message_text(
        text,
        parse_mode="Markdown",
        reply_markup=blacklist_menu_keyboard(),
    )


async def bl_remove_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    await query.edit_message_text(
        "❌ Введите *номер телефона* для удаления из чёрного списка:",
        parse_mode="Markdown",
    )
    return BL_REMOVE_PHONE


async def bl_remove_phone(update: Update, context: ContextTypes.DEFAULT_TYPE):
    phone = update.message.text.strip()
    removed = db.remove_from_blacklist(phone)
    keyboard = [[InlineKeyboardButton("⬅️ ЧС список", callback_data="blacklist_menu")]]
    if removed:
        text = f"✅ Номер *{phone}* удалён из чёрного списка."
    else:
        text = f"😕 Номер *{phone}* не найден в чёрном списке."
    await update.message.reply_text(
        text,
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )
    return ConversationHandler.END


async def bl_cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    await update.message.reply_text(
        "❌ Отменено.",
        reply_markup=blacklist_menu_keyboard(),
    )
    return ConversationHandler.END


async def export_excel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    await query.edit_message_text("⏳ Формирую Excel файл...")

    clients = db.get_all_clients()

    wb = Workbook()
    ws = wb.active
    ws.title = "Клиенты"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    center = Alignment(horizontal="center", vertical="center")

    headers = ["ID", "Имя", "Телефон", "Теги", "Заметка", "Дата добавления", "Кто добавил"]
    col_widths = [6, 30, 18, 35, 40, 18, 22]

    for col, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        ws.column_dimensions[cell.column_letter].width = width

    ws.row_dimensions[1].height = 20

    for row_idx, client in enumerate(clients, 2):
        cid, name, phone, email, notes, created_at, added_by_uid, added_by_uname = (*client, None, None)[:8]
        tags = db.get_client_tags(cid)
        tags_str = ", ".join(t[1] for t in tags) if tags else ""
        who = f"@{added_by_uname}" if added_by_uname else (f"#{added_by_uid}" if added_by_uid else "")
        values = [cid, name, phone, tags_str, notes or "", created_at[:10], who]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row_idx].height = 18

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"clients_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    keyboard = [[InlineKeyboardButton("⬅️ В меню", callback_data="menu")]]

    await query.message.reply_document(
        document=buffer,
        filename=filename,
        caption=f"📊 Клиентская база — *{len(clients)}* записей\n📅 {datetime.now().strftime('%d.%m.%Y %H:%M')}",
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )
    await query.edit_message_text(
        "✅ Файл Excel сформирован и отправлен.",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )


async def backup_db(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    await query.edit_message_text("⏳ Создаю резервную копию базы данных...")
    keyboard = [[InlineKeyboardButton("⬅️ В меню", callback_data="menu")]]
    filename = f"crm_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
    with open(DB_PATH, "rb") as f:
        await query.message.reply_document(
            document=f,
            filename=filename,
            caption=f"💾 *Резервная копия базы данных*\n📅 {datetime.now().strftime('%d.%m.%Y %H:%M')}",
            parse_mode="Markdown",
            reply_markup=InlineKeyboardMarkup(keyboard),
        )
    await query.edit_message_text(
        "✅ Резервная копия отправлена.",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )


# ─── Text import helpers ────────────────────────────────────────────────────

# Matches phone numbers in various international formats:
# +30698111111 | +30 698 333 3333 | 6982222222 | +7 999 123-45-67 | 8(999)123-45-67
_PHONE_RE = re.compile(r'\+?\d[\d\s\-\(\)\.]{5,18}\d')

# Letters that look like a person/company name (Cyrillic + Latin)
_NAME_CHARS = re.compile(r'[А-ЯЁа-яёA-Za-z]')


def _normalize_phone(raw: str) -> str:
    digits = re.sub(r'\D', '', raw)
    # Normalize Russian: leading 8 → 7
    if digits.startswith('8'):
        digits = '7' + digits[1:]
    # Russian 11-digit (7XXXXXXXXXX)
    if len(digits) == 11 and digits.startswith('7'):
        return '+' + digits
    # Russian 10-digit without country code
    if len(digits) == 10 and not raw.strip().startswith('+'):
        return '+7' + digits
    # International with explicit + (e.g. +30698111111)
    if raw.strip().startswith('+') and len(digits) >= 7:
        return '+' + digits
    # Any other number with enough digits
    if len(digits) >= 7:
        return digits
    return ''


def _check_phone(phone: str) -> str:
    """Returns 'blacklisted', 'exists', or 'new' — queries live SQLite DB."""
    if db.is_blacklisted(phone):
        return 'blacklisted'
    if db.search_by_phone(phone):
        return 'exists'
    return 'new'


def _extract_name_from_line(line: str, phone_match: re.Match) -> str:
    """Remove the phone number from the line; return remaining text as name."""
    before = line[:phone_match.start()].strip()
    after = line[phone_match.end():].strip()
    candidate = (before + ' ' + after).strip()
    # Keep only parts that contain at least one letter
    parts = [p for p in candidate.split() if _NAME_CHARS.search(p)]
    return ' '.join(parts)


def _parse_contacts_from_text(text: str) -> list[tuple[str, str]]:
    """Return list of (normalized_phone, name) pairs."""
    lines = [ln.strip() for ln in text.splitlines()]
    contacts: list[tuple[str, str]] = []
    seen: set[str] = set()

    for i, line in enumerate(lines):
        if not line:
            continue
        for m in _PHONE_RE.finditer(line):
            raw = m.group()
            norm = _normalize_phone(raw)
            if not norm or norm in seen:
                continue
            seen.add(norm)

            # 1) Try name from same line (words around the phone)
            name = _extract_name_from_line(line, m)

            # 2) Fall back to adjacent lines that have no phone numbers
            if not name:
                for offset in (-1, 1, -2, 2):
                    idx = i + offset
                    if 0 <= idx < len(lines):
                        adj = lines[idx]
                        if adj and not _PHONE_RE.search(adj) and _NAME_CHARS.search(adj):
                            name = adj.strip()
                            break

            contacts.append((norm, name or "Без имени"))

    return contacts


async def import_text_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    await query.edit_message_text(
        "📋 Отправьте скопированный текст из Telegram.\n"
        "Я сам найду номера и названия.\n\n"
        "Пример:\n"
        "Анна салон +30698111111\n"
        "Макс доставка 6982222222\n"
        "Игорь\n"
        "+30 698 333 3333\n\n"
        "/cancel — отменить",
    )
    return IMPORT_TEXT


async def receive_import_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    keyboard = [[InlineKeyboardButton("⬅️ В меню", callback_data="menu")]]
    text = update.message.text or ""

    contacts = _parse_contacts_from_text(text)

    if not contacts:
        await update.message.reply_text(
            "❌ Не распознано: номера не найдены в тексте.\n\n"
            "Убедитесь, что в тексте есть номера телефонов.",
            reply_markup=InlineKeyboardMarkup(keyboard),
        )
        return ConversationHandler.END

    added = 0
    duplicates = 0
    blacklisted = 0
    failed = 0
    user = update.message.from_user

    for phone, name in contacts:
        try:
            status = _check_phone(phone)
            if status == 'blacklisted':
                blacklisted += 1
            elif status == 'exists':
                duplicates += 1
            else:
                db.add_client(
                    name=name,
                    phone=phone,
                    notes="Импортирован из текста",
                    added_by_user_id=user.id,
                    added_by_username=user.username,
                )
                added += 1
        except Exception as e:
            logger.error("Text import failed for %s: %s", phone, e)
            failed += 1

    result_lines = [f"✅ Добавлено: {added}", f"⚠️ Уже были: {duplicates}"]
    if blacklisted:
        result_lines.append(f"🚫 В чёрном списке: {blacklisted}")
    if failed:
        result_lines.append(f"❌ Не распознано: {failed}")

    await update.message.reply_text(
        "\n".join(result_lines),
        reply_markup=InlineKeyboardMarkup(keyboard),
    )
    return ConversationHandler.END


def main():
    token = os.environ.get("TELEGRAM_BOT_TOKEN")
    if not token:
        raise RuntimeError("TELEGRAM_BOT_TOKEN environment variable not set")

    app = Application.builder().token(token).build()

    add_conv = ConversationHandler(
        entry_points=[CallbackQueryHandler(add_client_start, pattern="^add_client$")],
        states={
            ADD_NAME: [MessageHandler(filters.TEXT & ~filters.COMMAND, add_name)],
            ADD_PHONE: [MessageHandler(filters.TEXT & ~filters.COMMAND, add_phone)],
            ADD_NOTES: [MessageHandler(filters.TEXT & ~filters.COMMAND, add_notes)],
            ADD_TAGS: [
                CallbackQueryHandler(toggle_add_tag, pattern=r"^addtag_\d+$"),
                CallbackQueryHandler(finish_add_tags, pattern="^addtags_done$"),
            ],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        per_message=False,
        per_chat=True,
    )

    search_name_conv = ConversationHandler(
        entry_points=[CallbackQueryHandler(search_by_name_start, pattern="^search_by_name$")],
        states={
            SEARCH_BY_NAME: [MessageHandler(filters.TEXT & ~filters.COMMAND, do_search_by_name)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        per_message=False,
        per_chat=True,
    )

    search_phone_conv = ConversationHandler(
        entry_points=[CallbackQueryHandler(search_by_phone_start, pattern="^search_by_phone$")],
        states={
            SEARCH_BY_PHONE: [MessageHandler(filters.TEXT & ~filters.COMMAND, do_search_by_phone)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        per_message=False,
        per_chat=True,
    )

    edit_conv = ConversationHandler(
        entry_points=[CallbackQueryHandler(edit_field_prompt, pattern=r"^editfield_\d+_\w+$")],
        states={
            EDIT_VALUE: [MessageHandler(filters.TEXT & ~filters.COMMAND, edit_value_save)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        per_message=False,
        per_chat=True,
    )

    bl_add_conv = ConversationHandler(
        entry_points=[CallbackQueryHandler(bl_add_start, pattern="^bl_add$")],
        states={
            BL_ADD_PHONE: [MessageHandler(filters.TEXT & ~filters.COMMAND, bl_add_phone)],
            BL_ADD_REASON: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, bl_add_reason),
                CommandHandler("skip", bl_skip_reason),
            ],
        },
        fallbacks=[CommandHandler("cancel", bl_cancel)],
        per_message=False,
        per_chat=True,
    )

    bl_remove_conv = ConversationHandler(
        entry_points=[CallbackQueryHandler(bl_remove_start, pattern="^bl_remove$")],
        states={
            BL_REMOVE_PHONE: [MessageHandler(filters.TEXT & ~filters.COMMAND, bl_remove_phone)],
        },
        fallbacks=[CommandHandler("cancel", bl_cancel)],
        per_message=False,
        per_chat=True,
    )

    app.add_handler(MessageHandler(filters.ALL, _private_block_message), group=-1)
    app.add_handler(CallbackQueryHandler(_private_block_callback), group=-1)

    app.add_handler(CommandHandler("start", start))
    app.add_handler(add_conv)
    app.add_handler(search_name_conv)
    app.add_handler(search_phone_conv)
    app.add_handler(edit_conv)
    app.add_handler(bl_add_conv)
    app.add_handler(bl_remove_conv)

    import_text_conv = ConversationHandler(
        entry_points=[CallbackQueryHandler(import_text_start, pattern="^import_text$")],
        states={
            IMPORT_TEXT: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_import_text)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        per_message=False,
        per_chat=True,
    )
    app.add_handler(import_text_conv)

    app.add_handler(CallbackQueryHandler(menu, pattern="^menu$"))
    app.add_handler(CallbackQueryHandler(find_client_menu, pattern="^find_client$"))
    app.add_handler(CallbackQueryHandler(all_clients, pattern="^all_clients$"))
    app.add_handler(CallbackQueryHandler(delete_client, pattern=r"^delete_\d+$"))
    app.add_handler(CallbackQueryHandler(confirm_delete, pattern=r"^confirm_delete_\d+$"))
    app.add_handler(CallbackQueryHandler(edit_client, pattern=r"^edit_\d+$"))
    app.add_handler(CallbackQueryHandler(blacklist_menu, pattern="^blacklist_menu$"))
    app.add_handler(CallbackQueryHandler(bl_show, pattern="^bl_show$"))
    app.add_handler(CallbackQueryHandler(export_excel, pattern="^export_excel$"))
    app.add_handler(CallbackQueryHandler(backup_db, pattern="^backup_db$"))
    app.add_handler(CallbackQueryHandler(search_by_tag_menu, pattern="^search_by_tag$"))
    app.add_handler(CallbackQueryHandler(do_search_by_tag, pattern=r"^stag_\d+$"))
    app.add_handler(CallbackQueryHandler(client_tags_menu, pattern=r"^client_tags_\d+$"))
    app.add_handler(CallbackQueryHandler(toggle_client_tag, pattern=r"^ctag_\d+_\d+$"))
    app.add_handler(CallbackQueryHandler(finish_client_tags, pattern=r"^ctags_done_\d+$"))

    logger.info("Бот запускается...")
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
